#!/usr/bin/env python3
"""Import, render, and validate the splash-ad defense catalog."""

from __future__ import annotations

import argparse
import concurrent.futures
import datetime as dt
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

import openpyxl
import yaml


ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "data" / "tools.yml"
PROJECT_TITLE_EN = "Splash Ads Countermeasures"
PROJECT_TITLE_ZH = "开屏广告应对措施"
README_FILE = ROOT / "README.md"
CATALOG_FILE = ROOT / "CATALOG.md"
ENTRIES_DIR = ROOT / "entries"
README_ZH_FILE = ROOT / "README.zh-CN.md"
CATALOG_ZH_FILE = ROOT / "CATALOG.zh-CN.md"

EXPECTED_HEADERS = [
    "ID",
    "名称",
    "URL",
    "简介",
    "工作流程/原理",
    "工具类型",
    "主要防护对象",
    "局限性",
    "条目层级",
    "平台",
    "使用门槛",
    "状态",
]
STATUSES = {"可用", "维护情况不明", "已停更/归档", "不可获取", "待核验"}
EFFORT_LEVELS = {"low", "moderate", "high", "developer", "unverified"}
PUBLIC_EFFORT_LABELS = {
    "low": "🟢 低门槛",
    "moderate": "🟡 中等门槛",
    "high": "🔴 高门槛",
    "developer": "🧑‍💻 仅限开发者/平台方",
    "unverified": "❓ 待核验",
}
SECTION_ORDER = ("A.1", "A.2", "A.3", "A.4")
SECTION_TITLES = {
    "A.1": "广告处理工具与规则",
    "A.2": "应用和广告平台提供的设置",
    "A.3": "暂不纳入主目录的条目",
    "A.4": "系统与浏览器自带能力",
}
SECTION_EXPECTED_COUNTS = {"A.1": 110, "A.2": 15, "A.3": 17, "A.4": 4}
CATALOG_SECTION_ORDER = ("tools", "settings", "other", "limited", "cooperative", "general")
CATALOG_SECTION_TITLES = {
    "tools": "🛠️ 开屏相关工具、规则与研究原型",
    "settings": "📱 用户可开启的厂商广告控制",
    "other": "🧩 其他广告形式的防护工具",
    "limited": "🧪 实现或可获取性仍待核实的候选条目",
    "cooperative": "🏢 需要应用开发者或广告平台配合的机制",
    "general": "⚙️ 与广告防护相关的通用平台能力",
}
CATALOG_SECTION_TITLES_EN = {
    "tools": "🛠️ Splash-related tools, rules, and prototypes",
    "settings": "📱 User-enabled vendor advertising controls",
    "other": "🧩 Defense tools for other ad formats",
    "limited": "🧪 Candidates with unverified implementations or limited availability",
    "cooperative": "🏢 Mechanisms requiring host-app developer or advertising-platform action",
    "general": "⚙️ General platform capabilities related to ad defense",
}
CATALOG_SECTION_EXPECTED_COUNTS = {
    "tools": 88,
    "settings": 5,
    "other": 15,
    "limited": 22,
    "cooperative": 12,
    "general": 4,
}
REQUIRED_FIELDS = {
    "id",
    "name",
    "url",
    "description",
    "platform",
    "access",
    "effort",
    "status",
    "section",
    "catalog_section",
}
WORKFLOW_LABELS = ("触发/感知入口：", "识别/判定依据：", "干预点与手段：", "最终效果：")
MARKDOWN_LINK_RE = re.compile(r"(?<!!)\[[^\]]+\]\(([^)]+)\)")
HTML_HREF_RE = re.compile(r"<a\b[^>]*\bhref=[\"']([^\"']+)[\"']", re.IGNORECASE)
PUBLIC_STATUS_LABELS = {
    "可用": "已确认部署路径",
    "维护情况不明": "维护/兼容性不明确",
    "已停更/归档": "已停止维护",
    "不可获取": "当前无法获取",
    "待核验": "待核验",
}
PUBLIC_STATUS_EXPLANATIONS = {
    "已确认部署路径": "已找到可识别的安装包、服务、配置或源码构建路径；这不等于已验证当前兼容性或实际防护效果。",
    "维护/兼容性不明确": "维护状态或当前版本兼容性不清楚，使用前需要核对来源和版本。",
    "已停止维护": "项目已经停更或归档；历史机制仍可记录，但不代表当前版本仍可使用。",
    "当前无法获取": "目前无法从记录的来源获得可用版本或完整材料。",
    "待核验": "现有公开资料不足以确认实现方式、部署路径或当前可用性。",
}
PUBLIC_STATUS_LABELS_EN = {
    "可用": "Deployment path confirmed",
    "维护情况不明": "Maintenance / compatibility unclear",
    "已停更/归档": "Maintenance ended",
    "不可获取": "Currently unavailable",
    "待核验": "Needs verification",
}
PUBLIC_STATUS_EXPLANATIONS_EN = {
    "Deployment path confirmed": "An identifiable installation package, service, configuration, or source-build path has been found; this does not mean current compatibility or actual protection effectiveness has been verified.",
    "Maintenance / compatibility unclear": "The maintenance status or current-version compatibility is unclear; check the source and version before use.",
    "Maintenance ended": "The project has stopped maintenance or been archived; the historical mechanism may still be documented, but this does not mean it remains usable on current versions.",
    "Currently unavailable": "A usable version or complete material cannot currently be obtained from the recorded source.",
    "Needs verification": "The available public information is insufficient to confirm the implementation, deployment path, or current availability.",
}
PUBLIC_EFFORT_LABELS_EN = {
    "low": "🟢 Low effort",
    "moderate": "🟡 Moderate effort",
    "high": "🔴 High effort",
    "developer": "Developer/platform only",
    "unverified": "Not enough information",
}
PUBLIC_TEXT_REPLACEMENTS = (
    ("Resource Blocking / Modification", "广告请求或资源过滤"),
    ("Initialization / Startup Prevention", "启动阶段处理"),
    ("Presentation Suppression / Dismissal", "界面隐藏或自动关闭"),
    ("Input / Trigger Restriction", "输入或触发限制"),
    ("Navigation Interception / Confirmation", "跳转拦截或二次确认"),
    ("Mechanism awaiting verification or limited scope", "实现方式仍待核实，或适用范围有限"),
    ("Author-listed functions concern:", "资料中提到的功能包括："),
    ("general capability", "通用能力"),
    ("Source only / build needed.", "只有源码，需要自行构建。"),
    ("No prebuilt installation package was found", "未找到可直接安装的预编译包"),
    ("implementation awaiting verification", "实现方式仍待核实"),
    ("source code awaiting verification", "源码仍待核实"),
    ("implementation unverified", "实现方式未核实"),
    ("historical implementation", "历史版本实现"),
    ("historical source", "历史源码"),
    ("Developer configuration", "开发者配置"),
    ("Developer setting", "开发者设置"),
    ("User setting", "用户设置"),
    ("listed in the official guide", "记录在官方指南中"),
    ("用户设置 listed in the official guide", "记录在官方指南中的用户设置"),
    ("models and firmware supporting the option", "支持该选项的机型和固件"),
    ("corresponding HarmonyOS versions", "对应的 HarmonyOS 版本"),
    ("cited guide is version", "引用指南版本为"),
    ("release notes", "版本说明"),
    ("Official notes for", "官方说明适用于"),
    ("author-provided runtime environment", "作者提供的运行环境"),
    ("platform and package identity unverified", "平台和包身份尚未核实"),
    ("Android distribution package", "Android 分发包"),
    ("feature-specific authorization or module loading", "特定功能授权或模块加载"),
    ("some plugins also require application re-signing", "部分插件还需要对应用重新签名"),
    ("Xposed or privileges required by each feature", "Xposed 或各功能所需的相应权限"),
    ("or a corresponding loading environment", "或相应的加载环境"),
    ("host integration of the library, subscription configuration, and JavaScript support", "将该库集成到宿主应用、配置订阅并启用 JavaScript 支持"),
    ("developers configure entitlements and filtering services", "开发者配置权限和过滤服务"),
    ("Browser implementation, system settings, and user choices", "浏览器实现、系统设置和用户选择"),
    ("Developer sets mediation request parameters", "开发者设置聚合请求参数"),
    ("Developer supplies privacy-control configuration", "开发者提供隐私控制配置"),
    ("Developer sets the corresponding interface to NO", "开发者将对应接口设置为 NO"),
    ("Developer setting per ad request", "开发者按每次广告请求设置"),
    ("Developer setting before the request", "开发者在请求前设置"),
    ("Developer initialization configuration", "开发者初始化配置"),
    ("some ad sources have additional version requirements", "部分广告来源还有额外版本要求"),
    ("some sources require newer versions", "部分来源需要更高版本"),
    ("Manufacturer system support and user setting", "厂商系统支持和用户设置"),
    ("Accessibility Service", "无障碍服务"),
    ("accessibility service", "无障碍服务"),
    ("local VPN configuration", "本地 VPN 配置"),
    ("compatible proxy-rule clients", "兼容的代理规则客户端"),
    ("Paid.", "需要付费。"),
    ("Rules / scripts are free", "规则/脚本本身免费"),
    ("Only the chosen compatible client is needed.", "只需要安装所选的兼容客户端。"),
    ("Current product is browser-only", "当前产品仅适用于浏览器"),
    ("browser-only", "仅适用于浏览器"),
    ("The guide provides no complete model and firmware compatibility matrix.", "官方指南没有提供完整的机型和固件兼容列表。"),
    ("Documentation does not list complete ad-source coverage or default values.", "文档没有列出完整的广告来源覆盖范围或默认值。"),
    ("Provide interface-matching and skip-action rules", "提供界面匹配和跳过操作规则"),
    ("Use advertising terms, countdowns, and playback state to close ads or swipe up", "根据广告文字、倒计时和播放状态关闭广告或向上滑动"),
    ("Write advertising-domain and connection rules to Samsung network controls", "向 Samsung 网络控制功能写入广告域名和连接规则"),
    ("Provide domain rules for advertising servers and libraries", "提供广告服务器和广告库的域名规则"),
    ("Remove splash-ad fields and change display duration, dimensions, and expiration", "移除开屏广告字段，并修改展示时长、尺寸和过期时间"),
    ("Send domain queries to the author's configured remote ad-filtering service", "将域名查询发送到作者配置的远程广告过滤服务"),
    ("Filter advertising domains through a cloud resolution service", "通过云端解析服务过滤广告域名"),
    ("Remove matching advertising components and change the startup entry to the home screen", "移除匹配的广告组件，并将启动入口改为主界面"),
    ("Add confirmation before entering an ad destination", "在进入广告目标页面前增加一次确认"),
    ("Globally stop delivery of the corresponding sensor events", "全局停止传递对应的传感器事件"),
    ("Deny motion and orientation access during application startup", "在应用启动期间拒绝运动和方向传感器访问"),
    ("The system allows or denies requests using an address dataset", "系统根据地址数据集允许或拒绝网络请求"),
    ("Browsers conditionally confirm, allow, or deny external-app launches; the system routes links by association", "浏览器可以对打开外部应用进行确认、允许或拒绝，系统再根据关联关系处理链接"),
    ("Locates LinkedIn click methods and installs hooks, with the aim of handling in-app ads.", "定位 LinkedIn 的点击方法并安装 Hook，目标是处理应用内广告。"),
    ("Author describes direct home-screen entry, hidden ads, and intercepted ad loading", "资料描述了直接进入主界面、隐藏广告和拦截广告加载等功能"),
    ("Author describes intercepting specified QQ splash-cache use and ad navigation", "资料描述了拦截 QQ 指定的开屏缓存使用和广告跳转"),
    ("Author lists page replacement, view blocking, and request and file rules", "资料列出了页面替换、界面屏蔽以及请求和文件规则"),
    ("Author lists ad handling in Douyin, Kuaishou, and other applications, including some startup ads", "资料列出了抖音、快手等应用中的广告处理功能，其中包括部分开屏广告"),
    ("Provide configuration entries for advertising plugins across applications and mini programs", "为多个应用和小程序提供广告插件配置项"),
    ("Ad-skipping application listed in distribution records", "分发记录中列出的广告跳过应用"),
    ("Application in distribution records; advertising operations unverified", "分发记录中的应用，但广告相关功能尚未核实"),
    ("Download records and contemporaneous reports indicate ad-skipping use", "下载记录和同期资料显示其曾被用于跳过广告"),
    ("Add a reminder for detected accidental ad navigation", "检测到疑似误触跳转时增加提醒"),
    ("Trigger evaluation: configures splash-ad shaking through setSplashShakeButton", "通过 setSplashShakeButton 配置开屏广告是否响应摇一摇"),
    ("Input acquisition: denies access to motion or orientation data per app", "按应用限制运动或方向传感器的访问"),
    ("Input acquisition: disables SDK sensor use through setSensorStatus(false)", "通过 setSensorStatus(false) 禁止 SDK 使用传感器"),
    ("Input acquisition: controls SDK sensor access through isCanUseSensor", "通过 isCanUseSensor 控制 SDK 是否可以访问传感器"),
    ("Trigger evaluation: disables motion interaction components through isCanUseMotionManager", "通过 isCanUseMotionManager 禁用运动交互组件"),
    ("Trigger evaluation: configures permitted or prohibited shake, rotation, and swipe interactions", "配置是否允许摇一摇、旋转和滑动等交互"),
    ("Trigger evaluation: configures motion interaction features of ad sources", "配置广告来源的运动交互功能"),
    ("Input acquisition: passes the sensor collection choice to supported ad sources", "把传感器采集开关传递给支持的广告来源"),
    ("Input acquisition: restricts motion data; some firmware supports blocking only during splash ads", "限制运动数据访问；部分固件只支持在开屏广告期间阻止"),
    ("Trigger evaluation: disables splash-ad shaking with shakable=``0''", "通过 shakable=``0'' 禁用开屏广告摇一摇"),
    ("Input acquisition: restricts access to orientation data per app", "按应用限制方向传感器访问"),
    ("Input acquisition: restricts motion and orientation access per app", "按应用限制运动和方向传感器访问"),
    ("Trigger evaluation: disables splash-ad shaking through disableShake", "通过 disableShake 禁用开屏广告摇一摇"),
    ("Enabled by Google for selected placements", "由 Google 针对特定广告位启用"),
    ("Limited to Google Ads demand; specific splash-format coverage requires evaluation.", "仅适用于 Google Ads 广告需求；对具体开屏广告形式的覆盖范围仍需评估。"),
    ("Not blocked by default; affects only the specified interaction.", "默认不会阻止；只影响指定的交互方式。"),
    ("Normal orientation functions in the application are also restricted.", "应用本身的正常方向功能也会受到限制。"),
    ("General startup permission; official documentation does not limit it to advertising control.", "这是通用的启动权限，官方资料没有将其限定为广告控制功能。"),
    ("The interface contract covers interaction components; other advertising interactions may remain.", "该接口只覆盖交互组件，其他广告交互仍可能存在。"),
    ("Version evidence applies to the listed models; normal motion-based functions may be affected.", "版本证据只适用于列出的机型，正常的运动相关功能也可能受到影响。"),
    ("Omission or a value of 1 does not block shaking; documentation does not uniquely identify the parameter's first release.", "省略该参数或设置为 1 都不会阻止摇一摇；文档也无法明确该参数首次出现的版本。"),
    ("Scope depends on accidental-activation detection and the navigation being checked.", "覆盖范围取决于系统能否识别误触，以及对应的跳转是否经过检查。"),
    ("Some configurations are inaccessible; the Tencent Video plugin declares discontinued maintenance and version restrictions.", "部分配置无法获取；腾讯视频插件已声明停止维护并存在版本限制。"),
    ("Package identity and implementation are unconfirmed; download channels were previously closed.", "包身份和实现方式尚未确认，相关下载渠道此前已经关闭。"),
    ("Original documentation site is unavailable; correspondence between the package and feature descriptions remains unverified.", "原始说明网站已无法访问，安装包与功能描述之间的对应关系仍未核实。"),
    ("A renaming relationship with ``Chanshi'' has not been established, so its implementation description cannot be directly adopted.", "尚未确认它与“Chanshi”的更名关系，因此不能直接采用后者的实现描述。"),
    ("Advertising use is primarily supported by historical reports; current advertising implementation remains unverified.", "广告用途主要来自历史资料，当前广告实现仍未核实。"),
    ("Public distribution tree lacks implementation; specific intervention operations remain unverified.", "公开分发目录缺少实现代码，具体处理操作仍未核实。"),
    ("The inspected before / after callbacks are empty and do not yet remove ads. Not counted as an implemented defense, and display-control capabilities are not inferred from the project's goals.", "检查到的前后回调目前为空，尚未实现广告移除；因此不计作已实现的防护，也不会仅根据项目目标推断其具备展示控制能力。"),
    ("Corresponding public source is unavailable; the author explicitly excludes some real-time splash and feed ads.", "对应的公开源码不可获取；作者还明确排除了部分实时开屏和信息流广告。"),
    ("The corresponding advertising functionality lacks implementation source code", "对应广告功能缺少实现源码"),
    ("File replacement is documented in distribution materials; source evidence and version support vary by feature.", "分发资料记录了文件替换功能，但不同功能的源码证据和版本支持情况不同。"),
    ("Internal modification mechanism is undisclosed; listed as a candidate pending implementation verification.", "内部修改机制未公开，因此仅作为等待实现核验的候选线索。"),
    ("Resource Supply: Blocks Xuexitong splash-ad requests by URL.", "广告资源处理：按 URL 拦截学习通的开屏广告请求。"),
    ("Current product is browser-only; separate from the historical AdClear native-application filtering product.", "当前产品仅适用于浏览器，与历史上的 AdClear 原生应用过滤产品不同。"),
    ("Public tree lacks the application implementation and described build workflow; listed as an unverified candidate.", "公开代码目录缺少应用实现和所述构建流程，因此只能作为未核实线索。"),
    ("The inspected before / after callbacks are empty and do not yet remove ads.", "检查到的前后回调目前为空，尚未实现广告移除。"),
    ("Corresponding public source is unavailable", "对应的公开源码不可获取"),
    ("The corresponding advertising functionality lacks implementation source code", "对应广告功能缺少实现源码"),
    ("File replacement is documented in distribution materials", "分发资料中记录了文件替换功能"),
    ("Internal modification mechanism is undisclosed", "内部修改机制未公开"),
    ("Applies only to", "仅适用于"),
    ("Depends on", "取决于"),
    ("requires", "需要"),
    ("not counted", "不计入"),
    ("not verified", "未核实"),
)


class CatalogError(ValueError):
    pass


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def public_status(status: str, lang: str = "zh-CN") -> str:
    if lang == "en":
        return PUBLIC_STATUS_LABELS_EN.get(clean(status), "Not enough information")
    return PUBLIC_STATUS_LABELS.get(clean(status), "信息不足")


def effort_summary(tool: dict[str, Any], lang: str = "zh-CN") -> str:
    if lang == "en":
        translated = clean((tool.get("en") or {}).get("effort_display"))
        if translated:
            return translated
    efforts = tool.get("effort") or []
    labels = PUBLIC_EFFORT_LABELS_EN if lang == "en" else PUBLIC_EFFORT_LABELS
    separator = "; " if lang == "en" else "；"
    if len(efforts) == 1 and not clean(efforts[0].get("path")):
        return labels[efforts[0]["level"]]
    path_labels = {
        "内置或本地规则": "Built-in or local rules",
        "导入订阅规则": "Imported subscription rules",
        "原版内置规则": "Built-in rules (original release)",
        "手动导入社区规则": "Manual community-rule import",
        "内置节点、文字和图像识别": "Built-in node, text, and image recognition",
        "额外下载并配置 VLM": "Additional VLM download and setup",
        "本地 VPN": "Local VPN",
        "Root hosts": "Root hosts",
        "VPN": "VPN",
        "Root": "Root",
        "无障碍自动跳过": "Accessibility auto-skip",
        "无障碍 + Shizuku 点击": "Accessibility + Shizuku click",
        "Root + Xposed/LSPosed": "Root + Xposed/LSPosed",
        "本地 VPN/DNS": "Local VPN/DNS",
        "Shizuku": "Shizuku",
        "修改后的 APK": "Modified APK",
        "Xposed/LSPosed": "Xposed/LSPosed",
        "本地 VPN/DNS 过滤": "Local VPN/DNS filtering",
        "HTTPS 内容过滤和证书配置": "HTTPS content filtering and certificate setup",
        "VPN/DNS 解析器": "VPN/DNS resolver",
        "Root 网络控制": "Root network control",
        "Root 代理/iptables": "Root proxy/iptables",
        "导入 VPN/域名过滤规则": "Import VPN/domain filtering rules",
        "Root hosts 规则": "Root hosts rules",
        "安装项目提供的修改版 APK": "Install the project's modified APK",
        "自行修改、重签名和安装 APK": "Modify, re-sign, and install an APK",
        "导入 Loon 插件规则": "Import Loon plugin rules",
        "需要应用重签名的插件": "Plugin requiring app re-signing",
    }
    def path_label(value: str) -> str:
        if lang != "en":
            return value
        translated = path_labels.get(value, value)
        for source, target in {
            "代理": "proxy",
            "修改后的": "Modified ",
            "内容过滤和证书配置": "content filtering and certificate setup",
            "规则": "rules",
            "点击": "click",
        }.items():
            translated = translated.replace(source, target)
        return translated
    return separator.join(
        f"{path_label(clean(item.get('path')))}: {labels[item['level']]}"
        for item in efforts
    )


def localized(tool: dict[str, Any], field: str, lang: str) -> str:
    if lang == "en":
        value = (tool.get("en") or {}).get(field)
        if value:
            return clean(value)
    return clean(tool.get(field))


def localized_platform(tool: dict[str, Any], lang: str) -> str:
    return localized(tool, "platform", lang) or ("To verify" if lang == "en" else "待核验")


def language_switch(lang: str, english_target: str, chinese_target: str) -> str:
    """Render a compact, button-like language switch for GitHub Markdown."""
    if lang == "en":
        return f'<a href="{chinese_target}"><kbd>中文</kbd></a> <kbd>English</kbd>'
    return f'<kbd>中文</kbd> <a href="{english_target}"><kbd>English</kbd></a>'


def access_summary_en(tool: dict[str, Any]) -> str:
    return localized(tool, "access", "en") or "Deployment requirements need verification"

def has_cjk(text: str) -> bool:
    return bool(re.search(r"[\u3400-\u9fff]", text))


def looks_like_english(text: str) -> bool:
    letters = re.findall(r"[A-Za-z]", text)
    cjk = re.findall(r"[\u3400-\u9fff]", text)
    return len(letters) >= 24 and len(letters) > len(cjk) * 2


def public_text(value: Any, fallback: str = "") -> str:
    text = clean(value)
    if not text:
        return fallback or "待补充。"
    for source, target in PUBLIC_TEXT_REPLACEMENTS:
        text = text.replace(source, target)
    text = text.replace("用户设置 listed in the official guide", "记录在官方指南中的用户设置")
    text = text.replace("用户设置 记录在官方指南中", "记录在官方指南中的用户设置")
    text = re.sub(r"\bTrigger evaluation:\s*", "触发判断：", text)
    text = re.sub(r"\bInput acquisition:\s*", "输入获取：", text)
    text = re.sub(r"\bResource Supply:\s*", "广告资源处理：", text)
    text = re.sub(r"\bDisplay and Dismissal:\s*", "界面展示与关闭：", text)
    text = re.sub(r"\bInitialization and Startup:\s*", "初始化与启动：", text)
    text = re.sub(r"\bThe system\b", "系统", text)
    text = re.sub(r"\bThe author\b", "作者", text)
    text = re.sub(r"\bAuthor describes\b", "资料描述了", text)
    text = re.sub(r"\bAuthor lists\b", "资料列出了", text)
    text = re.sub(r"\bPublic tree lacks\b", "公开代码目录缺少", text)
    return text


def access_summary(tool: dict[str, Any], lang: str = "zh-CN") -> str:
    if lang == "en":
        return access_summary_en(tool)
    access = public_text(tool.get("access"))
    lowered = access.casefold()
    platform = public_text(tool.get("platform")).casefold()
    labels: list[str] = []

    has_root = any(token in lowered for token in ("root", "magisk"))
    has_jailbreak = "jailbreak" in lowered or "越狱" in access
    if has_root or has_jailbreak:
        if "android" in platform and "ios" not in platform:
            labels.append("Root")
        elif "ios" in platform and "android" not in platform:
            labels.append("Jailbreak / 越狱")
        else:
            labels.append("Root / Jailbreak")
    if any(token in lowered for token in ("runtime injection", "xposed", "lsposed", "libxposed", "vector", "substrate", "code injection")) or "运行时注入" in access:
        labels.append("Runtime Injection / 运行时注入")
    if any(token in lowered for token in ("apk modification", "modified apk", "repackag", "apk patch")) or ("apk" in lowered and any(word in access for word in ("修改", "补丁", "重打包", "重签名"))):
        labels.append("APK Modification / APK 修改")
    if "accessibility" in lowered or "无障碍" in access:
        labels.append("Accessibility / 无障碍")
    if "debugging" in lowered or "shizuku" in lowered or "调试" in access:
        labels.append("Debugging / 调试" + ("（Shizuku）" if "shizuku" in lowered else ""))
    if any(token in lowered for token in ("vpn", "dns", "proxy", "代理")):
        labels.append("VPN / DNS / 代理")
    if "overlay" in lowered or "悬浮" in access:
        labels.append("Overlay / 悬浮层")
    if "notification" in lowered or "通知" in access:
        labels.append("Notification Access / 通知访问")
    if any(token in lowered for token in ("source build", "build tool", "compile", "source code", "源码")):
        labels.append("Source Build / 源码构建")
    if any(token in lowered for token in ("additional model", "vlm", "model download")) or "模型" in access:
        labels.append("Additional Model / 额外模型")
    if "system setting" in lowered or "系统设置" in access:
        labels.append("System Setting / 系统设置")
    if any(token in lowered for token in ("developer", "host-app", "integration", "开发者")):
        labels.append("开发者 / 平台配置")

    if labels:
        return "；".join(dict.fromkeys(labels))
    if access in {"待核验", "信息不足"}:
        return "部署要求仍待核实"
    if looks_like_english(access):
        return "需要按原始资料配置相应的平台或运行环境"
    return access

def safety_notice(tool: dict[str, Any]) -> str:
    access = public_text(tool.get("access"))
    lowered = access.casefold()
    notices: list[str] = []
    if any(token in lowered for token in ("root", "xposed", "lsposed", "magisk", "jailbreak", "重签名", "越狱")):
        notices.append("可能影响系统安全或应用稳定性")
    if "apk" in lowered and any(token in lowered for token in ("修改", "modification", "patch", "补丁", "重打包")):
        notices.append("修改版应用可能带来签名、更新、账号和供应链风险")
    if any(token in lowered for token in ("vpn", "dns", "proxy", "代理")):
        notices.append("会接管部分网络连接或域名解析")
    if "accessibility" in lowered or "无障碍" in access:
        notices.append("请只对可信来源授予无障碍权限")
    if any(token in lowered for token in ("source only", "build needed", "research prototype", "研究原型", "源码")):
        notices.append("更适合开发者或研究者，不能按普通 App 直接安装")
    if tool.get("section") == "A.2":
        notices.append("这是应用或广告平台设置，不是独立安装包")
    if tool.get("section") == "A.4":
        notices.append("这是系统或浏览器能力，不是独立安装包")
    if tool.get("section") == "A.3":
        notices.append("仅作为补充线索，不作为普通用户推荐")
    return "；".join(dict.fromkeys(notices)) or "使用前请核对来源、版本和权限。"


def friendly_description(tool: dict[str, Any]) -> str:
    description = public_text(tool.get("description"))
    if not looks_like_english(description):
        return description
    if tool.get("section") == "A.2":
        return "应用或广告平台提供的设置，用于减少特定广告触发、展示或跳转。"
    if tool.get("section") == "A.4":
        return "系统或浏览器提供的通用能力，可能帮助减少广告触发、跳转或资源加载。"
    if tool.get("section") == "A.3":
        return "公开资料曾描述其与广告处理有关，但目前不作为主目录工具。"
    return "用于处理广告请求、广告界面或广告跳转，具体范围取决于项目支持的应用和版本。"


def friendly_limitation(tool: dict[str, Any]) -> str:
    limitation = public_text(tool.get("limitations"))
    if not looks_like_english(limitation):
        return limitation
    if tool.get("section") == "A.3":
        return "现有公开资料不足以确认实现方式、版本范围或当前可用性。"
    return "效果取决于应用版本、规则覆盖范围和运行环境；使用前请核对原始来源。"


def friendly_workflow(tool: dict[str, Any]) -> str:
    workflow = public_text(tool.get("workflow"))
    if not looks_like_english(workflow):
        return workflow
    tool_type = public_text(tool.get("tool_type"))
    return f"公开资料显示，该条目主要通过“{tool_type}”处理广告相关请求、界面或跳转；具体实现仍应以原始来源为准。"


def friendly_type(tool: dict[str, Any]) -> str:
    value = public_text(tool.get("tool_type"))
    if value == "待补充。" or looks_like_english(value):
        lowered = value.casefold()
        if "resource" in lowered or "过滤" in value:
            return "广告请求或资源过滤"
        if "presentation" in lowered or "关闭" in value:
            return "界面隐藏或自动关闭"
        if "input" in lowered or "触发" in value:
            return "输入或触发限制"
        if "navigation" in lowered or "跳转" in value:
            return "跳转拦截或二次确认"
        return "实现方式仍待核实"
    return value


def friendly_target(tool: dict[str, Any]) -> str:
    target = public_text(tool.get("protection_target"))
    if target in {"附录所述广告场景", "补充线索；不计入主工具清单"}:
        return "与该工具支持的应用广告、启动页、网络请求或广告跳转有关的场景。"
    if target == "厂商或广告平台控制所涉及的广告触发、展示或导航行为":
        return "特定应用或广告平台控制的广告触发、展示和跳转。"
    if target == "可被应用于广告防御的系统或平台能力":
        return "系统或浏览器层面的通用能力，可能被用于减少广告干扰。"
    return target


def canonical_intervention_points(tool: dict[str, Any], lang: str = "en") -> str:
    """Map implementation wording to the catalog's six intervention points."""
    source = localized(tool, "tool_type", "en") or clean(tool.get("tool_type"))
    lowered = source.casefold()
    points: list[str] = []

    def add(en: str, zh: str) -> None:
        points.append(en if lang == "en" else zh)

    if any(token in lowered for token in ("resource", "network", "dns", "hosts", "response", "cache")):
        add("Blocking and Modifying Ad Resources", "阻断和修改广告资源")
    if any(token in lowered for token in ("initialization", "startup", "entry-point")):
        add("Preventing Ad Initialization and Startup", "阻止广告初始化与启动")
    if any(token in lowered for token in ("presentation", "dismissal", "ui hiding", "visible ui", "automatic dismissal", "interface hiding")):
        add("Suppressing Ad Presentation and Automatically Dismissing Ads", "抑制广告展示与自动关闭广告")
    if any(token in lowered for token in ("input", "trigger", "sensor", "touch control")):
        add("Restricting Inputs and Disabling Triggers", "限制输入与禁用触发机制")
    if any(token in lowered for token in ("navigation blocking", "navigation control", "navigation interception", "confirmation")):
        add("Intercepting Ad Navigation and Requiring Confirmation", "拦截广告跳转与要求二次确认")
    if any(token in lowered for token in ("post-redirect", "return after navigation", "recovery after navigation", "post-navigation recovery")):
        add("Automatically Returning After Navigation", "跳转后自动返回")

    if points:
        return "; ".join(dict.fromkeys(points))
    if tool.get("section") == "A.3":
        return "Unverified / limited mechanism" if lang == "en" else "实现或适用范围仍待核实"
    return source or ("Intervention point needs verification" if lang == "en" else "干预位置仍待核实")

def safety_notice_en(tool: dict[str, Any]) -> str:
    access = localized(tool, "access", "en").casefold()
    notices: list[str] = []
    if any(token in access for token in ("root", "xposed", "lsposed", "magisk", "jailbreak", "re-sign", "injection")):
        notices.append("May affect system security or app stability")
    if any(token in access for token in ("apk modification", "modified apk", "repackag", "patch")):
        notices.append("Modified packages can introduce signing, update, account, and supply-chain risks")
    if any(token in access for token in ("vpn", "dns", "proxy")):
        notices.append("May take over part of network traffic or name resolution")
    if "accessibility" in access:
        notices.append("Grant Accessibility only to a source you trust")
    if any(token in access for token in ("source build", "research", "build tool")):
        notices.append("Better suited to developers or researchers; not directly installable")
    if tool.get("section") == "A.2":
        notices.append("Platform setting, not a standalone installable app")
    if tool.get("section") == "A.4":
        notices.append("System or browser capability, not a standalone installable app")
    if tool.get("section") == "A.3":
        notices.append("Supplementary lead, not a recommendation")
    result = "; ".join(dict.fromkeys(notices)) or "Check the source, version, and permissions before use."
    return result if result.endswith((".", "!", "?")) else result + "."


def detail_text_en(value: Any) -> str:
    return clean(value) or "Not provided."

def normalize_url(url: str) -> str:
    parsed = urllib.parse.urlsplit(url.strip())
    path = parsed.path.rstrip("/") or "/"
    return urllib.parse.urlunsplit(
        (parsed.scheme.lower(), parsed.netloc.lower(), path, parsed.query, "")
    )


def url_from_cell(cell: openpyxl.cell.cell.Cell) -> str:
    hyperlink = cell.hyperlink
    if hyperlink:
        target = hyperlink.target if hasattr(hyperlink, "target") else str(hyperlink)
        if target:
            return target.strip()

    value = clean(cell.value)
    if value.startswith(("http://", "https://")):
        return value

    match = re.match(r'^=HYPERLINK\("([^"]+)"', value, flags=re.IGNORECASE)
    return match.group(1).strip() if match else ""


def load_catalog(path: Path = DATA_FILE) -> dict[str, Any]:
    try:
        payload = yaml.safe_load(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise CatalogError(f"数据文件不存在: {path}") from exc
    except yaml.YAMLError as exc:
        raise CatalogError(f"YAML 无法解析: {exc}") from exc
    if not isinstance(payload, dict) or not isinstance(payload.get("tools"), list):
        raise CatalogError("data/tools.yml 必须包含 tools 列表")
    return payload


def validate_catalog(payload: dict[str, Any]) -> list[dict[str, Any]]:
    tools = payload["tools"]
    expected_count = payload.get("meta", {}).get("item_count")
    if expected_count != len(tools):
        raise CatalogError(f"meta.item_count={expected_count!r}，实际条目数为 {len(tools)}")

    ids: set[int] = set()
    names: set[str] = set()
    urls: set[str] = set()
    section_counts = {section: 0 for section in SECTION_ORDER}
    catalog_section_counts = {section: 0 for section in CATALOG_SECTION_ORDER}
    errors: list[str] = []

    for index, tool in enumerate(tools, start=1):
        if not isinstance(tool, dict):
            errors.append(f"第 {index} 项不是对象")
            continue
        missing = sorted(field for field in REQUIRED_FIELDS if not tool.get(field))
        if missing:
            errors.append(f"第 {index} 项缺少字段: {', '.join(missing)}")
            continue
        if "en" in (payload.get("meta", {}).get("languages") or []):
            required_en = {"description", "workflow", "tool_type", "protection_target", "limitations", "tier", "platform", "access"}
            en = tool.get("en")
            if not isinstance(en, dict) or any(not clean(en.get(field)) for field in required_en):
                errors.append(f"ID {tool.get('id')} 缺少完整英文展示字段")

        section = clean(tool["section"])
        if section not in SECTION_ORDER:
            errors.append(f"ID {tool.get('id')} 附录分区无效: {section!r}")
        else:
            section_counts[section] += 1

        catalog_section = clean(tool["catalog_section"])
        if catalog_section not in CATALOG_SECTION_ORDER:
            errors.append(f"ID {tool.get('id')} 总表分类无效: {catalog_section!r}")
        else:
            catalog_section_counts[catalog_section] += 1

        tool_id = tool["id"]
        if not isinstance(tool_id, int) or tool_id <= 0:
            errors.append(f"第 {index} 项 ID 非正整数: {tool_id!r}")
        elif tool_id in ids:
            errors.append(f"重复 ID: {tool_id}")
        ids.add(tool_id)

        name_key = clean(tool["name"]).casefold()
        if name_key in names:
            errors.append(f"重复名称: {tool['name']}")
        names.add(name_key)

        url = clean(tool["url"])
        parsed = urllib.parse.urlsplit(url)
        if parsed.scheme not in {"http", "https"} or not parsed.netloc:
            errors.append(f"ID {tool_id} URL 无效: {url!r}")
        url_key = normalize_url(url)
        if url_key in urls:
            errors.append(f"重复 URL: {url}")
        urls.add(url_key)

        if tool["status"] not in STATUSES:
            errors.append(f"ID {tool_id} 状态无效: {tool['status']!r}")

        efforts = tool.get("effort")
        if not isinstance(efforts, list) or not efforts:
            errors.append(f"ID {tool_id} 部署难度必须是非空列表")
        else:
            effort_paths: set[str] = set()
            for effort_index, effort in enumerate(efforts, start=1):
                if not isinstance(effort, dict):
                    errors.append(f"ID {tool_id} 第 {effort_index} 个部署路径不是对象")
                    continue
                level = clean(effort.get("level"))
                path = clean(effort.get("path"))
                if level not in EFFORT_LEVELS:
                    errors.append(f"ID {tool_id} 部署难度无效: {level!r}")
                if len(efforts) > 1 and not path:
                    errors.append(f"ID {tool_id} 有多个部署路径时必须填写路径名称")
                if path and path in effort_paths:
                    errors.append(f"ID {tool_id} 部署路径重复: {path}")
                effort_paths.add(path)

    ordered_ids = [tool.get("id") for tool in tools if isinstance(tool, dict)]
    if ordered_ids != sorted(ordered_ids):
        errors.append("条目没有按 ID 升序排列")
    declared_counts = payload.get("meta", {}).get("section_counts")
    if declared_counts is not None and declared_counts != section_counts:
        errors.append(
            f"meta.section_counts={declared_counts!r}，实际分区数量为 {section_counts!r}"
        )
    if declared_counts is not None and declared_counts != SECTION_EXPECTED_COUNTS:
        errors.append(
            f"附录分区数量必须为 {SECTION_EXPECTED_COUNTS!r}，实际声明为 {declared_counts!r}"
        )
    declared_catalog_counts = payload.get("meta", {}).get("catalog_section_counts")
    if declared_catalog_counts is not None and declared_catalog_counts != catalog_section_counts:
        errors.append(
            f"meta.catalog_section_counts={declared_catalog_counts!r}，实际总表分类数量为 {catalog_section_counts!r}"
        )
    if declared_catalog_counts is not None and declared_catalog_counts != CATALOG_SECTION_EXPECTED_COUNTS:
        errors.append(
            f"总表分类数量必须为 {CATALOG_SECTION_EXPECTED_COUNTS!r}，实际声明为 {declared_catalog_counts!r}"
        )
    if errors:
        raise CatalogError("\n".join(errors))
    return tools


def import_xlsx(
    source: Path,
    sheet_name: str,
    output: Path,
    expected_count: int,
    section: str = "A.1",
) -> None:
    workbook = openpyxl.load_workbook(source, data_only=False, read_only=False)
    if sheet_name not in workbook.sheetnames:
        raise CatalogError(f"找不到工作表 {sheet_name!r}；现有: {', '.join(workbook.sheetnames)}")
    sheet = workbook[sheet_name]
    headers = [clean(sheet.cell(1, column).value) for column in range(1, 13)]
    if headers != EXPECTED_HEADERS:
        raise CatalogError(
            "表头不匹配\n"
            f"期望: {' | '.join(EXPECTED_HEADERS)}\n"
            f"实际: {' | '.join(headers)}"
        )

    tools: list[dict[str, Any]] = []
    legacy_catalog_sections = {
        "A.1": "tools",
        "A.2": "settings",
        "A.3": "limited",
        "A.4": "general",
    }
    for row_number in range(2, sheet.max_row + 1):
        raw_id = sheet.cell(row_number, 1).value
        if raw_id in (None, ""):
            continue
        try:
            tool_id = int(raw_id)
        except (TypeError, ValueError) as exc:
            raise CatalogError(f"第 {row_number} 行 ID 无效: {raw_id!r}") from exc

        values = [clean(sheet.cell(row_number, column).value) for column in range(2, 13)]
        name, _url_display, description, workflow, tool_type, target, limitations, tier, platform, access, status = values
        tools.append(
            {
                "id": tool_id,
                "name": name,
                "url": url_from_cell(sheet.cell(row_number, 3)),
                "description": description,
                "workflow": workflow,
                "tool_type": tool_type,
                "protection_target": target,
                "limitations": limitations,
                "tier": tier,
                "platform": platform or "待核验",
                "access": access or "待核验",
                "effort": [{"level": "unverified"}],
                "status": status or "待核验",
                "section": section,
                "catalog_section": legacy_catalog_sections.get(section, "tools"),
            }
        )

    if len(tools) != expected_count:
        raise CatalogError(f"期望 {expected_count} 条，实际导入 {len(tools)} 条")

    payload = {
        "meta": {
            "schema_version": 1,
            "title": PROJECT_TITLE_EN,
            "title_zh": PROJECT_TITLE_ZH,
            "language": "zh-CN",
            "verified_on": dt.date.today().isoformat(),
            "item_count": len(tools),
        },
        "tools": tools,
    }
    validate_catalog(payload)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        yaml.safe_dump(payload, allow_unicode=True, sort_keys=False, width=1000),
        encoding="utf-8",
    )


def table_text(value: Any) -> str:
    return clean(value).replace("|", "\\|").replace("\r", " ").replace("\n", " ")


def detail_text(value: Any) -> str:
    text = public_text(value)
    if not text:
        return "待补充。"
    lines = []
    for line in text.splitlines():
        rendered = line
        for label in WORKFLOW_LABELS:
            if line.startswith(label):
                public_label = {
                    "触发/感知入口：": "从哪里开始：",
                    "识别/判定依据：": "如何判断：",
                    "干预点与手段：": "怎么处理：",
                    "最终效果：": "可能结果：",
                }[label]
                rendered = f"**{public_label}**{line[len(label):]}"
                break
        if looks_like_english(rendered):
            if "如何判断" in rendered:
                rendered = "**如何判断：**依据项目规则、文档或已知的应用界面判断目标广告相关内容。"
            elif "怎么处理" in rendered:
                rendered = "**怎么处理：**按照项目说明在请求、界面或跳转阶段进行处理。"
            elif "可能结果" in rendered:
                rendered = "**可能结果：**具体效果和覆盖范围仍取决于项目实现、应用版本和运行环境。"
            elif "从哪里开始" in rendered:
                rendered = "**从哪里开始：**需要按原始项目说明配置相应的平台、权限或运行环境。"
            else:
                rendered = "现有公开资料为英文，具体含义请以原始来源为准。"
        lines.append(rendered)
    return "<br>\n".join(lines)


def render_full_catalog(payload: dict[str, Any], lang: str = "zh-CN") -> str:
    meta = payload["meta"]
    titles = CATALOG_SECTION_TITLES_EN if lang == "en" else CATALOG_SECTION_TITLES
    section_counts = meta.get("catalog_section_counts") or {
        section: sum(1 for tool in payload["tools"] if tool.get("catalog_section") == section)
        for section in CATALOG_SECTION_ORDER
    }
    section_count_text = ("; " if lang == "en" else "；").join(
        (f"{titles[section]} {section_counts.get(section, 0)}"
         if lang == "en" else f"{titles[section]} {section_counts.get(section, 0)} 条")
        for section in CATALOG_SECTION_ORDER
    )
    rows = [
        ("# Complete catalog" if lang == "en" else "# 完整目录"),
        "",
        language_switch(lang, "CATALOG.md", "CATALOG.zh-CN.md"),
        "",
        (f"This catalog comprises {meta['item_count']} records: 108 primary countermeasure entries (93 related to splash ads and 15 targeting other ad formats), 22 candidates with limited or unverified evidence, 12 mechanisms requiring developer or platform support, and 4 general platform capabilities."
         if lang == "en" else f"本目录共记录 {meta['item_count']} 个条目：其中 108 个是主要防护条目（93 个与开屏广告相关，15 个针对其他广告形式），另有 22 个证据或可获取性有限的候选条目、12 个需要开发者或广告平台配合的机制，以及 4 个通用平台能力。"),
        "",
        "> [!IMPORTANT]",
        ("> Inclusion does not imply recommendation. Tool effectiveness, compatibility, and maintenance status can change across versions; for entries involving Accessibility, VPN/DNS, Root, Xposed, re-signing, or experimental code, verify the source, permissions, and risks first."
         if lang == "en" else "> 收录不代表推荐。工具效果、兼容性和维护状态可能随版本变化；涉及无障碍权限、VPN/DNS、Root、Xposed、重签名或实验性代码的条目，请先确认来源、权限和风险。"),
        "",
        (f"- Records: {meta['item_count']}" if lang == "en" else f"- 条目数：{meta['item_count']}"),
        (f"- Section counts: {section_count_text}" if lang == "en" else f"- 分类数量：{section_count_text}"),
        (f"- Verification baseline: {meta['verified_on']}" if lang == "en" else f"- 内容核验基线：{meta['verified_on']}"),
        ("- Purpose: provide checkable tool information; this is not installation, purchasing, or security advice" if lang == "en" else "- 本目录用途：提供可核对的工具信息，不构成安装、购买或安全建议"),
        ("- Vendor settings, developer-cooperative mechanisms, and system capabilities are not necessarily standalone tools that ordinary users can install independently" if lang == "en" else "- 厂商设置、开发者配合机制和系统能力不一定是普通用户可以独立安装的工具"),
        ("- Candidate entries are supplementary leads only and are not included in ordinary-user tool recommendations" if lang == "en" else "- 候选条目只作为补充线索，不纳入普通用户工具推荐"),
        "",
    ]
    column_headers = {
        "tools": (("Defense", "Platform", "Deployment effort", "Deployment requirements", "What it does", "Availability / limits") if lang == "en" else ("方案", "平台", "部署难度", "部署要求", "能做什么", "可获取性/限制")),
        "settings": (("Setting / control", "Platform or scope", "Deployment effort", "Deployment requirements", "What it does", "Limits") if lang == "en" else ("设置/控制项", "适用平台或范围", "部署难度", "部署要求", "能做什么", "限制")),
        "other": (("Defense", "Platform", "Deployment effort", "Deployment requirements", "What it does", "Availability / limits") if lang == "en" else ("方案", "平台", "部署难度", "部署要求", "能做什么", "可获取性/限制")),
        "limited": (("Entry", "Platform", "Deployment effort", "What it does", "Why it is not included in the main inventory", "Availability") if lang == "en" else ("条目", "平台", "部署难度", "能做什么", "暂不纳入主清单的原因", "可获取性")),
        "cooperative": (("Mechanism / control", "Platform or scope", "Deployment effort", "Deployment requirements", "What it does", "Limits") if lang == "en" else ("机制/控制项", "适用平台或范围", "部署难度", "部署要求", "能做什么", "限制")),
        "general": (("System / browser capability", "Platform", "Deployment effort", "Deployment requirements", "What it does", "Limits") if lang == "en" else ("系统/浏览器能力", "平台", "部署难度", "部署要求", "能做什么", "局限")),
    }
    for section in CATALOG_SECTION_ORDER:
        rows.extend([f"## {titles[section]}", ""])
        headers = column_headers[section]
        rows.append("| " + " | ".join(headers) + " |")
        rows.append("| " + " | ".join("---" for _ in headers) + " |")
        section_tools = [tool for tool in payload["tools"] if tool["catalog_section"] == section]
        for tool in section_tools:
            suffix = "" if lang == "en" else ".zh-CN"
            link = f"entries/{tool['id']:03d}{suffix}.md"
            if section in {"tools", "other"}:
                values = (
                    tool["name"],
                    localized_platform(tool, lang), effort_summary(tool, lang),
                    (access_summary_en(tool) if lang == "en" else access_summary(tool)),
                    localized(tool, "description", lang) if lang == "en" else friendly_description(tool),
                    f"{public_status(tool['status'], lang)}; {localized(tool, 'limitations', lang)}" if lang == "en" else f"{public_status(tool['status'])}；{friendly_limitation(tool)}",
                )
            elif section in {"settings", "cooperative"}:
                values = (
                    tool["name"],
                    localized_platform(tool, lang), effort_summary(tool, lang),
                    (access_summary_en(tool) if lang == "en" else access_summary(tool)),
                    localized(tool, "description", lang) if lang == "en" else friendly_description(tool),
                    localized(tool, "limitations", lang) if lang == "en" else friendly_limitation(tool),
                )
            elif section == "limited":
                values = (
                    tool["name"],
                    localized_platform(tool, lang), effort_summary(tool, lang),
                    localized(tool, "description", lang) if lang == "en" else friendly_description(tool),
                    localized(tool, "limitations", lang) if lang == "en" else friendly_limitation(tool),
                    public_status(tool["status"], lang),
                )
            else:
                values = (
                    tool["name"],
                    localized_platform(tool, lang), effort_summary(tool, lang),
                    (access_summary_en(tool) if lang == "en" else access_summary(tool)),
                    localized(tool, "description", lang) if lang == "en" else friendly_description(tool),
                    localized(tool, "limitations", lang) if lang == "en" else friendly_limitation(tool),
                )
            rendered = [table_text(value) for value in values]
            rendered[0] = f"[{rendered[0]}]({link})"
            rows.append("| " + " | ".join(rendered) + " |")
        rows.append("")
    rows.extend(
        [
            ("## 🧭 Deployment effort" if lang == "en" else "## 🧭 部署难度（Deployment Effort）"),
            "",
            ("- **🟢 Low effort**: install and authorize a ready-to-use app/service, or enable an existing system setting." if lang == "en" else "- **🟢 低门槛**：安装并授权现成的应用/服务即可使用，或只需开启已有系统设置。"),
            ("- **🟡 Moderate effort**: requires additional configuration such as rules, certificates, debugging pairing, region selection, or replacement-app setup." if lang == "en" else "- **🟡 中等门槛**：需要额外配置，例如导入规则、选择区域、安装证书、完成调试配对或安装修改版应用。"),
            ("- **🔴 High effort**: requires Root / jailbreak, runtime injection, APK modification, or source compilation." if lang == "en" else "- **🔴 高门槛**：需要 Root/越狱、运行时注入、APK 修改或源码编译。"),
            ("- **Developer/platform only**: requires SDK integration, host-app changes, or platform enablement; ordinary users cannot deploy it independently." if lang == "en" else "- **仅限开发者/平台方**：需要集成 SDK、修改宿主应用或由广告平台启用，不是普通用户可以独立部署的工具。"),
            ("- **Not enough information**: the available public information is insufficient to confirm an actionable deployment path." if lang == "en" else "- **信息不足**：现有公开资料不足以确认一条可操作的部署路径。"),
            "",
            ("Deployment effort is assessed per deployment path. Payment does not change the level, and effort is separate from safety, availability, and tested effectiveness." if lang == "en" else "部署难度按具体部署路径判断。同一条目如果有多种路径，会分别标注；是否付费不改变难度等级，部署难度也与安全性、可获取性和实测效果分开记录。"),
            "",
            ("## 🏷️ Availability and maintenance status" if lang == "en" else "## 🏷️ 可获取性与维护状态"),
            "",
            *[f"- **{label}**: {explanation}" for label, explanation in (PUBLIC_STATUS_EXPLANATIONS_EN.items() if lang == "en" else PUBLIC_STATUS_EXPLANATIONS.items())],
            "",
            ("## Deployment requirement notes" if lang == "en" else "## 使用门槛提示"),
            "",
            ("- **Root / Jailbreak**: a system privilege condition; this is separate from runtime injection." if lang == "en" else "- **Root / Jailbreak（越狱）**：系统权限条件，与运行时注入是不同的部署要求。"),
            ("- **Runtime Injection**: frameworks such as Xposed / LSPosed modify app behavior while it runs; this is not the same as APK modification." if lang == "en" else "- **Runtime Injection / 运行时注入**：例如 Xposed / LSPosed，在应用运行时改变其行为；它不等同于 APK 修改。"),
            ("- **APK Modification**: changes the installation package and requires a compatible patching, signing, and installation path." if lang == "en" else "- **APK Modification / APK 修改**：修改安装包，并需要兼容的补丁、签名和重新安装流程。"),
            ("- **Accessibility**: can expose interface information and allow automated actions; authorize only trusted sources." if lang == "en" else "- **Accessibility / 无障碍**：可能读取界面信息并执行自动操作，请只对可信来源授权。"),
            ("- **Debugging**: some paths use Shizuku or another debugging-authorized interface; this is separate from Accessibility." if lang == "en" else "- **Debugging / 调试**：部分路径通过 Shizuku 等方式获得调试授权，与无障碍权限分开。"),
            ("- **VPN / DNS / proxy**: may mediate part of network traffic or name resolution; HTTPS rewriting may additionally require certificate trust." if lang == "en" else "- **VPN / DNS / 代理**：可能接管部分网络连接或域名解析；HTTPS 改写还可能需要证书信任。"),
            ("- **Source Build**: requires compiling the defense tool itself; it is distinct from patching a target APK." if lang == "en" else "- **Source Build / 源码构建**：需要编译防护工具本身，与修改目标 APK 是不同要求。"),
            "",
        ]
    )
    rows.extend(
        [
            "",
            ("## Maintenance and licensing" if lang == "en" else "## 维护与许可"),
            "",
            ("Catalog data and explanatory text are licensed under [CC BY 4.0](LICENSE-CONTENT); generation and validation scripts use [MIT](LICENSE). See [CONTRIBUTING.md](CONTRIBUTING.md) for corrections." if lang == "en" else "目录数据和说明文字采用 [CC BY 4.0](LICENSE-CONTENT)，生成与校验脚本采用 [MIT](LICENSE)。纠错线索请参阅 [CONTRIBUTING.zh-CN.md](CONTRIBUTING.zh-CN.md)。"),
            "",
        ]
    )
    return "\n".join(rows)


def render_entry(tool: dict[str, Any], lang: str = "zh-CN") -> str:
    category_map = {
        "A.1": "Ad-handling tools and rules",
        "A.2": "Settings provided by apps and ad platforms",
        "A.3": "Entries not included in the main catalog",
        "A.4": "Built-in system and browser capabilities",
    }
    category = clean((tool.get("en") or {}).get("category")) if lang == "en" else SECTION_TITLES[tool["section"]]
    if lang == "en" and not category:
        category = category_map[tool["section"]]
    status = public_status(tool["status"], lang)
    name = tool["name"]
    platform = localized_platform(tool, lang)
    effort = effort_summary(tool, lang)
    requirements = access_summary_en(tool) if lang == "en" else access_summary(tool)
    description = localized(tool, "description", lang) if lang == "en" else friendly_description(tool)
    limitations = localized(tool, "limitations", lang) if lang == "en" else friendly_limitation(tool)
    workflow = localized(tool, "workflow", lang) if lang == "en" else friendly_workflow(tool)
    target = localized(tool, "protection_target", lang) if lang == "en" else friendly_target(tool)
    kind = localized(tool, "tool_type", lang) if lang == "en" else friendly_type(tool)
    if lang == "en":
        quick_labels = ("ID", "Inventory group", "Platform", "Deployment effort", "Deployment requirements", "Availability / maintenance")
        top_link = language_switch(lang, f"{tool['id']:03d}.md", f"{tool['id']:03d}.zh-CN.md")
        nav = f"[Back to guide](../README.md) · [Back to full catalog](../CATALOG.md) · [Original source]({tool['url']})"
        type_note = clean((tool.get("en") or {}).get("type_note")) or {
            "A.2": "App or ad-platform setting; not a standalone installable package",
            "A.4": "System or browser capability; not a standalone installable package",
            "A.3": "Supplementary lead; not an ordinary-user recommendation",
        }.get(tool["section"])
        rows = [f"| {quick_labels[0]} | {tool['id']:03d} |", f"| {quick_labels[1]} | {category} |", f"| {quick_labels[2]} | {table_text(platform)} |", f"| {quick_labels[3]} | {table_text(effort)} |", f"| {quick_labels[4]} | {table_text(requirements)} |", f"| {quick_labels[5]} | {table_text(status)} |"]
        if type_note:
            rows.append(f"| Type | {type_note} |")
        return "\n".join([
            f"# {name}", "", top_link, "", nav, "", "## Quick view", "", "| Field | Details |", "| --- | --- |", *rows, "", "### What it does", "", table_text(description), "", "### Limits and notes", "", table_text(limitations), "", "### Before use", "", table_text(clean((tool.get("en") or {}).get("safety_notice")) or safety_notice_en(tool)), "", "## Technical notes", "", "### Working method", "", detail_text_en(workflow), "", "### Coverage", "", detail_text_en(target), "", "### Intervention point", "", detail_text_en(kind), "", "## Original source", "", f"- {tool['url']}", "",])
    quick_rows = [
        f"| 编号 | {tool['id']:03d} |",
        f"| 收录分组 | {table_text(category)} |",
        f"| 平台 | {table_text(platform)} |",
        f"| 部署难度 | {table_text(effort)} |",
        f"| 部署要求 | {table_text(requirements)} |",
        f"| 可获取性 / 维护状态 | {table_text(status)} |",
    ]
    if tool["section"] == "A.2":
        quick_rows.append("| 类型 | 应用或广告平台设置，不是独立安装包 |")
    elif tool["section"] == "A.4":
        quick_rows.append("| 类型 | 系统或浏览器能力，不是独立安装包 |")
    elif tool["section"] == "A.3":
        quick_rows.append("| 类型 | 补充线索，不作为普通用户推荐 |")
    return "\n".join(
        [
            f"# {tool['name']}",
            "",
            language_switch(lang, f"{tool['id']:03d}.md", f"{tool['id']:03d}.zh-CN.md"),
            "",
            f"[返回使用导航](../README.zh-CN.md) · [返回完整目录](../CATALOG.zh-CN.md) · [原始来源]({tool['url']})",
            "",
            "## 快速了解",
            "",
            "| 字段 | 内容 |",
            "| --- | --- |",
            *quick_rows,
            "",
            "### 能做什么",
            "",
            table_text(description),
            "",
            "### 限制和注意事项",
            "",
            table_text(limitations),
            "",
            "### 使用前提示",
            "",
            table_text(safety_notice(tool)),
            "",
            "## 技术说明",
            "",
            "### 工作方式",
            "",
            detail_text(friendly_workflow(tool)),
            "",
            "### 可覆盖场景",
            "",
            detail_text(friendly_target(tool)),
            "",
            "### 干预位置",
            "",
            detail_text(friendly_type(tool)),
            "",
            "## 原始来源",
            "",
            f"- {tool['url']}",
            "",
        ]
    )


def expected_outputs(payload: dict[str, Any]) -> dict[Path, str]:
    outputs = {
        CATALOG_FILE: render_full_catalog(payload, "en"),
        CATALOG_ZH_FILE: render_full_catalog(payload, "zh-CN"),
    }
    for tool in payload["tools"]:
        outputs[ENTRIES_DIR / f"{tool['id']:03d}.md"] = render_entry(tool, "en")
        outputs[ENTRIES_DIR / f"{tool['id']:03d}.zh-CN.md"] = render_entry(tool, "zh-CN")
    return outputs


def render_catalog(payload: dict[str, Any]) -> None:
    outputs = expected_outputs(payload)
    ENTRIES_DIR.mkdir(parents=True, exist_ok=True)
    for path, content in outputs.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
    expected_entry_paths = {path for path in outputs if path.parent == ENTRIES_DIR}
    for stale in ENTRIES_DIR.glob("*.md"):
        if stale not in expected_entry_paths:
            stale.unlink()


def check_generated(payload: dict[str, Any]) -> None:
    problems = []
    for path, expected in expected_outputs(payload).items():
        if not path.exists():
            problems.append(f"缺少生成文件: {path.relative_to(ROOT)}")
        elif path.read_text(encoding="utf-8") != expected:
            problems.append(f"生成文件已过期: {path.relative_to(ROOT)}")
    actual_entries = set(ENTRIES_DIR.glob("*.md")) if ENTRIES_DIR.exists() else set()
    expected_entries = {p for p in expected_outputs(payload) if p.parent == ENTRIES_DIR}
    for path in sorted(actual_entries - expected_entries):
        problems.append(f"存在过期详情页: {path.relative_to(ROOT)}")
    if problems:
        raise CatalogError("\n".join(problems))


def check_catalog_links(tools: list[dict[str, Any]], root: Path = ROOT) -> None:
    problems = []
    def has_link(text: str, target: str) -> bool:
        return f"]({target})" in text or f'href="{target}"' in text or f"href='{target}'" in text
    readme = root / "README.md"
    readme_zh = root / "README.zh-CN.md"
    readme_text = readme.read_text(encoding="utf-8") if readme.exists() else ""
    readme_zh_text = readme_zh.read_text(encoding="utf-8") if readme_zh.exists() else ""
    catalog_file = root / "CATALOG.md"
    catalog_zh_file = root / "CATALOG.zh-CN.md"
    catalog_text = catalog_file.read_text(encoding="utf-8") if catalog_file.exists() else ""
    catalog_zh_text = catalog_zh_file.read_text(encoding="utf-8") if catalog_zh_file.exists() else ""
    if not has_link(readme_text, "CATALOG.md"):
        problems.append("README missing catalog link: CATALOG.md")
    if not has_link(readme_zh_text, "CATALOG.zh-CN.md"):
        problems.append("Chinese README missing catalog link")
    for tool in tools:
        for suffix, catalog_body, readme_name, catalog_name in (("", catalog_text, "README.md", "CATALOG.md"), (".zh-CN", catalog_zh_text, "README.zh-CN.md", "CATALOG.zh-CN.md")):
            entry_name = f"entries/{tool['id']:03d}{suffix}.md"
            if not has_link(catalog_body, entry_name):
                problems.append(f"Catalog missing entry link: {entry_name}")
            entry_path = root / entry_name
            if not entry_path.exists():
                continue
            entry_text = entry_path.read_text(encoding="utf-8")
            if f"../{readme_name}" not in entry_text:
                problems.append(f"Entry missing README link: {entry_name}")
            if f"../{catalog_name}" not in entry_text:
                problems.append(f"Entry missing catalog link: {entry_name}")
            if suffix == "" and not has_link(entry_text, f"{tool['id']:03d}.zh-CN.md"):
                problems.append(f"English entry missing language link: {entry_name}")
            if suffix == ".zh-CN" and not has_link(entry_text, f"{tool['id']:03d}.md"):
                problems.append(f"Chinese entry missing language link: {entry_name}")
            source_link = f"[Original source]({tool['url']})" if suffix == "" else f"[原始来源]({tool['url']})"
            if source_link not in entry_text:
                problems.append(f"Entry missing source link: {entry_name}")
    if problems:
        raise CatalogError("\n".join(problems))


def check_internal_links(root: Path = ROOT) -> None:
    problems = []
    for markdown_file in sorted(root.rglob("*.md")):
        if ".git" in markdown_file.parts:
            continue
        content = markdown_file.read_text(encoding="utf-8")
        targets = MARKDOWN_LINK_RE.findall(content) + HTML_HREF_RE.findall(content)
        for raw_target in targets:
            target = raw_target.strip().split(maxsplit=1)[0].strip("<>")
            parsed = urllib.parse.urlsplit(target)
            if parsed.scheme or target.startswith("#"):
                continue
            relative_path = urllib.parse.unquote(parsed.path)
            if not relative_path:
                continue
            resolved = (markdown_file.parent / relative_path).resolve()
            if not resolved.exists():
                source = markdown_file.relative_to(root)
                problems.append(f"Markdown 断链: {source} -> {target}")
    if problems:
        raise CatalogError("\n".join(problems))


def probe_url(url: str) -> tuple[str, str]:
    headers = {"User-Agent": "splashads-countermeasures-link-check/1.0"}
    for method in ("HEAD", "GET"):
        try:
            request = urllib.request.Request(url, headers=headers, method=method)
            with urllib.request.urlopen(request, timeout=10) as response:
                return url, str(response.status)
        except urllib.error.HTTPError as exc:
            if method == "HEAD" and exc.code in {403, 405}:
                continue
            return url, f"HTTP {exc.code}"
        except Exception as exc:  # Network reports must not change catalog state.
            if method == "HEAD":
                continue
            return url, f"{type(exc).__name__}: {exc}"
    return url, "unknown"


def report_external_links(tools: list[dict[str, Any]]) -> None:
    with concurrent.futures.ThreadPoolExecutor(max_workers=12) as executor:
        results = list(executor.map(probe_url, [tool["url"] for tool in tools]))
    failures = [(url, result) for url, result in results if not result.startswith("2")]
    print(f"外部链接检查：{len(results) - len(failures)}/{len(results)} 返回 2xx")
    for url, result in failures:
        print(f"WARN {result}: {url}")


def command_import(args: argparse.Namespace) -> None:
    import_xlsx(args.source, args.sheet, args.output, args.expected_count, args.section)
    payload = load_catalog(args.output)
    render_catalog(payload)
    print(f"已导入并生成 {len(payload['tools'])} 条记录")


def command_render(_args: argparse.Namespace) -> None:
    payload = load_catalog()
    validate_catalog(payload)
    render_catalog(payload)
    print(f"已生成中英文 CATALOG 和 {len(payload['tools']) * 2} 个详情页")


def command_check(args: argparse.Namespace) -> None:
    payload = load_catalog()
    tools = validate_catalog(payload)
    check_generated(payload)
    check_catalog_links(tools)
    check_internal_links()
    if args.external_links:
        report_external_links(tools)
    print(f"检查通过：{len(tools)} 条记录，生成文件一致")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    import_parser = subparsers.add_parser("import", help="从下载的工作表 .xlsx 导入单个分类")
    import_parser.add_argument("source", type=Path)
    import_parser.add_argument("--sheet", default="Tools")
    import_parser.add_argument("--output", type=Path, default=DATA_FILE)
    import_parser.add_argument("--expected-count", type=int, default=110)
    import_parser.add_argument("--section", choices=SECTION_ORDER, default="A.1")
    import_parser.set_defaults(func=command_import)

    render_parser = subparsers.add_parser("render", help="从 YAML 重新生成 Markdown")
    render_parser.set_defaults(func=command_render)

    check_parser = subparsers.add_parser("check", help="校验数据和生成文件")
    check_parser.add_argument("--external-links", action="store_true")
    check_parser.set_defaults(func=command_check)
    return parser


def main() -> int:
    try:
        args = build_parser().parse_args()
        args.func(args)
    except CatalogError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
