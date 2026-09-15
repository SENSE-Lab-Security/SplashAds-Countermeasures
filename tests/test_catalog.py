from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

import openpyxl


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "catalog.py"
SPEC = importlib.util.spec_from_file_location("catalog", SCRIPT)
catalog = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(catalog)


class CatalogTests(unittest.TestCase):
    def make_workbook(self, path: Path, count: int = 120) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Tools"
        sheet.append(catalog.EXPECTED_HEADERS)
        for tool_id in range(1, count + 1):
            sheet.append(
                [
                    tool_id,
                    f"Tool {tool_id}",
                    f"display {tool_id}",
                    f"description {tool_id}",
                    "触发/感知入口：测试。",
                    "test type",
                    "test target",
                    "test limitation",
                    "直接防护工具/客户端",
                    "",
                    "",
                    "",
                ]
            )
            sheet.cell(tool_id + 1, 3).hyperlink = f"https://example.com/tool/{tool_id}"
        workbook.save(path)

    def test_import_uses_hyperlink_and_conservative_defaults(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            output = Path(directory) / "tools.yml"
            self.make_workbook(source)
            catalog.import_xlsx(source, "Tools", output, 120, "A.1")
            payload = catalog.load_catalog(output)
            tools = catalog.validate_catalog(payload)
            self.assertEqual(120, len(tools))
            self.assertEqual("https://example.com/tool/1", tools[0]["url"])
            self.assertEqual("待核验", tools[0]["platform"])
            self.assertEqual("待核验", tools[0]["access"])
            self.assertEqual("待核验", tools[0]["status"])
            self.assertEqual("A.1", tools[0]["section"])
            self.assertEqual("tools", tools[0]["catalog_section"])

    def test_duplicate_url_is_rejected(self) -> None:
        payload = {
            "meta": {"item_count": 2},
            "tools": [
                {
                    "id": 1,
                    "name": "One",
                    "url": "https://example.com/tool",
                    "description": "one",
                    "platform": "Android",
                    "access": "Accessibility",
                    "effort": [{"level": "unverified"}],
                    "status": "待核验",
                    "section": "A.1",
                    "catalog_section": "tools",
                },
                {
                    "id": 2,
                    "name": "Two",
                    "url": "https://EXAMPLE.com/tool/",
                    "description": "two",
                    "platform": "Android",
                    "access": "Accessibility",
                    "effort": [{"level": "unverified"}],
                    "status": "待核验",
                    "section": "A.1",
                    "catalog_section": "tools",
                },
            ],
        }
        with self.assertRaisesRegex(catalog.CatalogError, "重复 URL"):
            catalog.validate_catalog(payload)

    def test_import_rejects_changed_headers(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            output = Path(directory) / "tools.yml"
            self.make_workbook(source)
            workbook = openpyxl.load_workbook(source)
            workbook["Tools"]["A1"] = "序号"
            workbook.save(source)
            with self.assertRaisesRegex(catalog.CatalogError, "表头不匹配"):
                catalog.import_xlsx(source, "Tools", output, 120)

    def test_internal_markdown_link_check_rejects_missing_target(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "README.md").write_text(
                "[存在](CONTRIBUTING.md)\n[断链](entries/999.md)\n",
                encoding="utf-8",
            )
            (root / "CONTRIBUTING.md").write_text("# OK\n", encoding="utf-8")
            with self.assertRaisesRegex(catalog.CatalogError, "Markdown 断链"):
                catalog.check_internal_links(root)

    def test_sparse_ids_are_allowed(self) -> None:
        payload = {
            "meta": {"item_count": 2},
            "tools": [
                {
                    "id": 1,
                    "name": "One",
                    "url": "https://example.com/one",
                    "description": "one",
                    "platform": "Android",
                    "access": "Accessibility",
                    "effort": [{"level": "unverified"}],
                    "status": "待核验",
                    "section": "A.1",
                    "catalog_section": "tools",
                },
                {
                    "id": 3,
                    "name": "Three",
                    "url": "https://example.com/three",
                    "description": "three",
                    "platform": "Android",
                    "access": "Accessibility",
                    "effort": [{"level": "unverified"}],
                    "status": "待核验",
                    "section": "A.1",
                    "catalog_section": "tools",
                },
            ],
        }
        self.assertEqual(2, len(catalog.validate_catalog(payload)))

    def test_canonical_catalog_has_all_appendix_sections(self) -> None:
        payload = catalog.load_catalog()
        tools = catalog.validate_catalog(payload)
        self.assertEqual(146, len(tools))
        self.assertEqual(3, payload["meta"]["schema_version"])
        self.assertEqual(
            {"A.1": 110, "A.2": 15, "A.3": 17, "A.4": 4},
            payload["meta"]["section_counts"],
        )
        self.assertEqual(
            {"tools": 88, "settings": 5, "other": 15, "limited": 22, "cooperative": 12, "general": 4},
            payload["meta"]["catalog_section_counts"],
        )
        self.assertTrue(all(tool["effort"] for tool in tools))

    def test_effort_paths_are_rendered_separately(self) -> None:
        payload = catalog.load_catalog()
        by_id = {tool["id"]: tool for tool in catalog.validate_catalog(payload)}
        self.assertEqual([{"level": "moderate"}], by_id[1]["effort"])
        self.assertEqual([{"level": "high"}], by_id[7]["effort"])
        self.assertEqual([{"level": "high"}], by_id[18]["effort"])
        self.assertEqual([{"level": "developer"}], by_id[140]["effort"])
        self.assertEqual([{"level": "unverified"}], by_id[155]["effort"])
        adaway = catalog.effort_summary(by_id[47])
        self.assertIn("本地 VPN", adaway)
        self.assertIn("🟢 低门槛", adaway)
        self.assertIn("Root", adaway)
        self.assertIn("🔴 高门槛", adaway)

    def test_full_catalog_uses_markdown_headings_for_sections(self) -> None:
        payload = catalog.load_catalog()
        full_catalog = catalog.render_full_catalog(payload)
        for title in catalog.CATALOG_SECTION_TITLES.values():
            self.assertIn(f"## {title}\n", full_catalog)
        for forbidden in ("Overleaf", "附录", "A.1", "A.2", "A.3", "A.4"):
            self.assertNotIn(forbidden, full_catalog)
        self.assertIn("## 🧭 部署难度（Deployment Effort）", full_catalog)
        self.assertIn("部署难度按具体部署路径判断", full_catalog)

    def test_readme_is_public_navigation(self) -> None:
        readme = catalog.README_FILE.read_text(encoding="utf-8")
        self.assertIn("Where should I start", readme)
        self.assertIn("](CATALOG.md)", readme)
        self.assertNotIn("## 暂不纳入主目录的条目", readme)

    def test_public_status_labels_are_rendered_without_changing_yaml(self) -> None:
        payload = catalog.load_catalog()
        full_catalog = catalog.render_full_catalog(payload)
        self.assertIn("已确认部署路径", full_catalog)
        self.assertIn("维护/兼容性不明确", full_catalog)
        self.assertIn("待核验", full_catalog)
        self.assertEqual("待核验", catalog.public_status("待核验"))
        self.assertTrue(any(tool["status"] == "待核验" for tool in payload["tools"]))

    def test_public_details_use_user_facing_sections(self) -> None:
        payload = catalog.load_catalog()
        by_id = {tool["id"]: tool for tool in payload["tools"]}
        ordinary = catalog.render_entry(by_id[1])
        excluded = catalog.render_entry(by_id[155])
        control = catalog.render_entry(by_id[140])
        self.assertIn("## 快速了解", ordinary)
        self.assertIn("## 技术说明", ordinary)
        self.assertIn("| 部署难度 |", ordinary)
        self.assertIn("暂不纳入主目录的条目", excluded)
        self.assertIn("不是独立安装包", control)
        self.assertNotIn("附录分区", ordinary)
        self.assertNotIn("条目层级", ordinary)

    def test_canonical_catalog_links_cover_entries_and_sources(self) -> None:
        payload = catalog.load_catalog()
        catalog.check_catalog_links(catalog.validate_catalog(payload))

    def test_bilingual_outputs_and_language_links(self) -> None:
        payload = catalog.load_catalog()
        english = catalog.render_full_catalog(payload, "en")
        chinese = catalog.render_full_catalog(payload, "zh-CN")
        self.assertIn('href="CATALOG.zh-CN.md"><kbd>中文</kbd>', english)
        self.assertIn('<kbd>中文</kbd> <a href="CATALOG.md"><kbd>English</kbd>', chinese)
        self.assertIn(f"## {catalog.CATALOG_SECTION_TITLES_EN['tools']}", english)
        self.assertIn(f"## {catalog.CATALOG_SECTION_TITLES['tools']}", chinese)
        self.assertIn("Deployment path confirmed", english)
        self.assertIn("已确认部署路径", chinese)
        entry = next(tool for tool in payload["tools"] if tool["id"] == 1)
        self.assertIn('href="001.zh-CN.md"><kbd>中文</kbd>', catalog.render_entry(entry, "en"))
        self.assertIn('<kbd>中文</kbd> <a href="001.md"><kbd>English</kbd>', catalog.render_entry(entry, "zh-CN"))


if __name__ == "__main__":
    unittest.main()
